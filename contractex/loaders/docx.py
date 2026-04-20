"""DOCX document loader using python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from contractex.exceptions import DocumentLoadError
from contractex.loaders.base import DocumentLoader


class DOCXLoader(DocumentLoader):
    """
    DOCX document loader using python-docx.

    Supports Microsoft Word .docx files.
    """

    def __init__(self, include_headers: bool = True, include_footers: bool = True):
        """
        Initialize DOCX loader.

        Args:
            include_headers: Include header text
            include_footers: Include footer text
        """
        self.include_headers = include_headers
        self.include_footers = include_footers

        # Check python-docx availability
        try:
            import docx

            self.docx = docx
        except ImportError as e:
            raise DocumentLoadError(
                "python-docx not installed. Install with: pip install python-docx"
            ) from e

    def load(self, source: str) -> str:
        """
        Load DOCX and extract text.

        Args:
            source: Path to DOCX file

        Returns:
            Extracted text content

        Raises:
            DocumentLoadError: If loading fails
        """
        try:
            path = Path(source)
            if not path.exists():
                raise FileNotFoundError(f"DOCX file not found: {source}")

            # Load document
            doc = self.docx.Document(source)  # type: ignore[attr-defined]

            text_parts = []

            # Extract headers
            if self.include_headers:
                for section in doc.sections:
                    header = section.header
                    for paragraph in header.paragraphs:
                        if paragraph.text.strip():
                            text_parts.append(paragraph.text)

            # Extract main content
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)

            # Extract footers
            if self.include_footers:
                for section in doc.sections:
                    footer = section.footer
                    for paragraph in footer.paragraphs:
                        if paragraph.text.strip():
                            text_parts.append(paragraph.text)

            # Combine all parts
            full_text = "\n".join(text_parts)

            return full_text

        except Exception as e:
            raise DocumentLoadError(f"Failed to load DOCX: {str(e)}") from e

    def get_metadata(self, source: str) -> dict[str, Any]:
        """Get DOCX metadata."""
        metadata = super().get_metadata(source)

        try:
            doc = self.docx.Document(source)  # type: ignore[attr-defined]

            # Add DOCX-specific metadata
            core_properties = doc.core_properties

            metadata.update(
                {
                    "author": core_properties.author,
                    "created": core_properties.created,
                    "modified": core_properties.modified,
                    "title": core_properties.title,
                    "subject": core_properties.subject,
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables),
                }
            )

        except Exception:
            pass

        return metadata

    def supports(self, file_path: str) -> bool:
        """Check if file is a DOCX."""
        return Path(file_path).suffix.lower() == ".docx"
